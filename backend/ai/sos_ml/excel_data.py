"""Task-specific Excel inputs. This source is embedded in both standalone notebooks."""
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from openpyxl import load_workbook
from .data import NEEDS, ENTITIES, fingerprint


def parse_marked_text(value):
    """Convert [LOCATION|xã A] into raw text and Unicode spans; no manual offsets."""
    value = unicodedata.normalize('NFC', str(value))
    pattern = re.compile(r'\[([A-Z_]+)\|([^\[\]]+)\]')
    chunks, entities, cursor, length = [], [], 0, 0
    for match in pattern.finditer(value):
        prefix = value[cursor:match.start()]
        if '[' in prefix or ']' in prefix:
            raise ValueError('Dấu nhãn không hợp lệ. Dùng [LOCATION|cụm từ], không lồng nhãn.')
        chunks.append(prefix)
        length += len(prefix)
        label, text = match.groups()
        if label not in ENTITIES or not text.strip() or text != text.strip():
            raise ValueError(f'Nhãn/cụm từ không hợp lệ: {match.group()}')
        entities.append({'start': length, 'end': length + len(text), 'label': label, 'text': text})
        chunks.append(text)
        length += len(text)
        cursor = match.end()
    tail = value[cursor:]
    if '[' in tail or ']' in tail:
        raise ValueError('Dấu nhãn chưa đóng hoặc sai cú pháp [NHAN|cụm từ].')
    chunks.append(tail)
    return ''.join(chunks), entities


def excel_bool(value):
    if value is True or str(value).strip().lower() in {'1', 'true', 'yes', 'có', 'co'}:
        return True
    if value is False or str(value).strip().lower() in {'0', 'false', 'no', 'không', 'khong'}:
        return False
    raise ValueError(f'Giá trị phải là 0 hoặc 1: {value!r}')


def load_excel_data(path, task, evaluation_domain='synthetic'):
    if task not in {'a', 'b'}:
        raise ValueError('task phải là a hoặc b')
    if evaluation_domain not in {'synthetic', 'real_vietnamese'}:
        raise ValueError('evaluation_domain không hợp lệ')
    workbook = load_workbook(path, read_only=True, data_only=False)
    required = ['id', 'group_id', 'text' if task == 'a' else 'text_annotated']
    required += ['is_request', 'needs'] if task == 'a' else []
    required += ['origin', 'reviewed', 'source']
    splits = {}
    seen_ids, seen_text, seen_groups = set(), {}, {}
    try:
        for sheet_name in ['Train', 'Validation', 'Test']:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f'Thiếu sheet {sheet_name}')
            iterator = workbook[sheet_name].iter_rows(values_only=True)
            headers = next(iterator, ())
            if len(set(headers)) != len(headers) or any(h not in headers for h in required):
                raise ValueError(f'{sheet_name}: cột bắt buộc: {required}. Không dùng file của model còn lại.')
            rows = []
            for number, values in enumerate(iterator, 2):
                if all(v is None or v == '' for v in values):
                    continue
                cells = dict(zip(headers, values))
                try:
                    if any(isinstance(v, str) and v.startswith('=') for v in values):
                        raise ValueError('Không dùng công thức Excel trong dữ liệu train.')
                    for column in required:
                        if column != 'needs' and (cells.get(column) is None or str(cells[column]).strip() == ''):
                            raise ValueError(f'Thiếu {column}')
                    if task == 'a':
                        text = unicodedata.normalize('NFC', str(cells['text']))
                        entities = []
                        request = excel_bool(cells['is_request'])
                        needs = [v.strip() for v in str(cells.get('needs') or '').split(';') if v.strip()]
                        if set(needs) - set(NEEDS) or len(needs) != len(set(needs)):
                            raise ValueError('needs chỉ dùng RESCUE;MEDICAL;EVACUATION;SUPPLY, không lặp nhãn.')
                        if not request and needs:
                            raise ValueError('is_request=0 phải để needs trống.')
                    else:
                        text, entities = parse_marked_text(cells['text_annotated'])
                        request, needs = bool(entities), []  # Internal collator placeholder; B does not learn request/needs.
                    if not text.strip():
                        raise ValueError('Nội dung trống')
                    origin = str(cells['origin']).strip()
                    reviewed = excel_bool(cells['reviewed'])
                    if origin not in {'synthetic', 'real', 'translated'}:
                        raise ValueError('origin phải là synthetic, real hoặc translated')
                    if origin != 'synthetic' and not reviewed:
                        raise ValueError('Dữ liệu thực/dịch cần được người gắn nhãn kiểm tra, rồi đặt reviewed=1.')
                    if sheet_name != 'Train':
                        expected = 'synthetic' if evaluation_domain == 'synthetic' else 'real'
                        if origin != expected:
                            raise ValueError(f'Chế độ đánh giá này yêu cầu origin={expected} cho Validation/Test.')
                    row = {'id': str(cells['id']).strip(), 'group_id': str(cells['group_id']).strip(),
                           'text': text, 'is_request': request, 'needs': needs, 'entities': entities,
                           'origin': origin, 'reviewed': reviewed, 'source_url': str(cells['source']).strip()}
                    key = fingerprint(text)
                    if row['id'] in seen_ids or key in seen_text:
                        raise ValueError(f'Trùng ID/nội dung (bản trước: {seen_text.get(key, row["id"])}).')
                    group = row['group_id']
                    if group in seen_groups and seen_groups[group] != sheet_name:
                        raise ValueError(f'group_id={group} xuất hiện ở nhiều tập: rò rỉ dữ liệu.')
                    seen_ids.add(row['id'])
                    seen_text[key] = row['id']
                    seen_groups[group] = sheet_name
                    rows.append(row)
                except (ValueError, TypeError) as exc:
                    raise ValueError(f'{sheet_name}, dòng Excel {number}: {exc}') from exc
            if not rows:
                raise ValueError(f'{sheet_name}: không có dữ liệu')
            if task == 'a':
                if {r['is_request'] for r in rows} != {True, False}:
                    raise ValueError(f'{sheet_name}: cần cả is_request=0 và 1')
                for need in NEEDS:
                    flags = [need in r['needs'] for r in rows]
                    if not any(flags) or all(flags):
                        raise ValueError(f'{sheet_name}: cần cả mẫu có và không có nhãn {need}')
            else:
                present = {e['label'] for r in rows for e in r['entities']}
                if set(ENTITIES) - present or all(r['entities'] for r in rows):
                    raise ValueError(f'{sheet_name}: cần đủ 4 loại thông tin và mẫu không có thông tin cần trích xuất')
            splits[sheet_name.lower()] = rows
    finally:
        workbook.close()
    return splits


def semantic_hash(rows, task):
    """Sorting Excel rows or changing its formatting does not change this fingerprint."""
    fields = ['id', 'group_id', 'text'] + (['is_request', 'needs'] if task == 'a' else ['entities'])
    records = [{k: row[k] for k in fields} for row in sorted(rows, key=lambda r: r['id'])]
    return hashlib.sha256(json.dumps(records, ensure_ascii=False, sort_keys=True).encode('utf-8')).hexdigest()


def check_excel_parent(parent, splits, task):
    old = parent.get('excel_split_sha256', {})
    for name in ['validation', 'test']:
        if old.get(name) != semantic_hash(splits[name], task):
            raise ValueError(f'{name}: không khớp dữ liệu đánh giá của model cha. Dùng lại Validation/Test của lần trước. '
                             'Học tiếp trong bộ Excel này yêu cầu model xuất từ notebook mới.')
