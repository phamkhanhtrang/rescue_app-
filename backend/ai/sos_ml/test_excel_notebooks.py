"""Offline checks: workbook semantics, leakage, continuation and notebook execution.

No real model is trained. Notebook helper cells are executed without Drive or GPU.
"""
import ast
import copy
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from .excel_data import load_excel_data, parse_marked_text, semantic_hash, check_excel_parent
from .preprocess import encode

ROOT = Path(__file__).resolve().parents[1] / 'training'


class ExcelTests(unittest.TestCase):
    def test_counts_and_separate_task_labels(self):
        a = load_excel_data(ROOT / 'data_model_a.xlsx', 'a')
        b = load_excel_data(ROOT / 'data_model_b.xlsx', 'b')
        for task, data, count in [('a', a, 960), ('b', b, 1032)]:
            self.assertEqual({k:len(v) for k,v in data.items()}, {'train':count,'validation':176,'test':160})
            self.assertTrue(all(not row['reviewed'] for rows in data.values() for row in rows))
        self.assertTrue(all(not r['entities'] for rows in a.values() for r in rows))
        self.assertTrue(all(not r['needs'] for rows in b.values() for r in rows))
        self.assertNotEqual({r['id'] for r in a['train']}, {r['id'] for r in b['train']})
        for name in ['validation', 'test']:
            self.assertEqual([r['text'] for r in a[name]], [r['text'] for r in b[name]])

    def test_wrong_workbook_and_false_real_claim_rejected(self):
        with self.assertRaisesRegex(ValueError, 'cột bắt buộc'):
            load_excel_data(ROOT / 'data_model_b.xlsx', 'a')
        with self.assertRaisesRegex(ValueError, 'origin=real'):
            load_excel_data(ROOT / 'data_model_a.xlsx', 'a', 'real_vietnamese')

    def test_annotations_repeated_phrase_unicode_and_invalid(self):
        text, spans = parse_marked_text('🆘 [LOCATION|xã A], không phải xã A, cần [RESOURCE|nước uống].')
        self.assertEqual(text[spans[0]['start']:spans[0]['end']], 'xã A')
        self.assertEqual(spans[0]['start'], 2)
        for value in ['[LOCATION|xã A', '[BOGUS|x]', '[LOCATION| xã A]', '[LOCATION|[RESOURCE|x]]']:
            with self.assertRaises(ValueError):
                parse_marked_text(value)

    def test_missing_label_and_leakage_errors_include_row(self):
        for column, value, message in [('D2', None, 'Thiếu is_request'), ('E2', 'WRONG', 'needs')]:
            wb = load_workbook(ROOT / 'data_model_a.xlsx')
            wb['Train'][column] = value
            with patch('ai.sos_ml.excel_data.load_workbook', return_value=wb):
                with self.assertRaisesRegex(ValueError, 'Train, dòng Excel 2:.*' + message):
                    load_excel_data('memory', 'a')
        wb = load_workbook(ROOT / 'data_model_a.xlsx')
        wb['Test']['C2'] = wb['Train']['C2'].value
        with patch('ai.sos_ml.excel_data.load_workbook', return_value=wb):
            with self.assertRaisesRegex(ValueError, 'Test, dòng Excel 2:.*Trùng'):
                load_excel_data('memory', 'a')

    def test_add_train_and_keep_holdouts_for_continue(self):
        splits = load_excel_data(ROOT / 'data_model_b.xlsx', 'b')
        parent = {'excel_split_sha256':{k:semantic_hash(v, 'b') for k,v in splits.items()}}
        extended = copy.deepcopy(splits)
        row = copy.deepcopy(extended['train'][0])
        row.update(id='new-row', group_id='new-family', text=row['text'] + ' Xin phản hồi.')
        extended['train'].append(row)
        check_excel_parent(parent, extended, 'b')
        extended['test'].reverse()
        check_excel_parent(parent, extended, 'b')
        extended['test'][0]['text'] += ' changed'
        with self.assertRaisesRegex(ValueError, 'test:'):
            check_excel_parent(parent, extended, 'b')

    def test_all_excel_rows_align_with_real_vietnamese_tokenizer(self):
        from transformers import AutoTokenizer
        tokenizer = AutoTokenizer.from_pretrained(ROOT.parent / 'nlp/sos_phobert_model', use_fast=False, local_files_only=True)
        for task in ['a','b']:
            splits = load_excel_data(ROOT / f'data_model_{task}.xlsx', task)
            for split, rows in splits.items():
                for row in rows:
                    with self.subTest(task=task, split=split, row=row['id']):
                        encode(row['text'], tokenizer, row['entities'] if task == 'b' else None)


class NotebookTests(unittest.TestCase):
    def test_failed_drive_mount_stops_before_creating_work_directory(self):
        import types
        from unittest.mock import Mock
        for task in ['a', 'b']:
            nb = json.loads((ROOT / f'0{1 if task == "a" else 2}_train_model_{task}.ipynb').read_text(encoding='utf-8'))
            source = next(''.join(c['source']) for c in nb['cells']
                          if c['cell_type'] == 'code' and 'drive.mount(' in ''.join(c['source']))
            colab = types.ModuleType('google.colab')
            original = ValueError('mount failed')
            colab.drive = types.SimpleNamespace(mount=Mock(side_effect=original))
            colab.files = Mock()
            with patch.dict('sys.modules', {'google.colab': colab}), patch('pathlib.Path.mkdir') as mkdir:
                with self.assertRaisesRegex(RuntimeError, 'Chưa kết nối được Google Drive') as caught:
                    exec(compile(source, '<colab-config>', 'exec'), {})
                self.assertIs(caught.exception.__cause__, original)
                mkdir.assert_not_called()

    def test_standalone_helper_cells_and_literal_colab_parameters(self):
        for task in ['a', 'b']:
            nb = json.loads((ROOT / f'0{1 if task == "a" else 2}_train_model_{task}.ipynb').read_text(encoding='utf-8'))
            namespace = {'__name__':'notebook_offline_check'}
            for cell in nb['cells']:
                if cell['cell_type'] != 'code':
                    continue
                source = ''.join(cell['source'])
                tree = ast.parse(source)
                compile(tree, '<notebook>', 'exec')
                self.assertFalse(any(isinstance(n, ast.ImportFrom) and n.level for n in ast.walk(tree)))
                if any(isinstance(n, (ast.ClassDef, ast.FunctionDef)) for n in tree.body):
                    exec(compile(tree, '<notebook_helpers>', 'exec'), namespace)
                for line in source.splitlines():
                    if '#@param' in line and 'type:' in line:
                        statement = ast.parse(line).body[0]
                        self.assertIsInstance(statement, ast.Assign)
                        self.assertIsInstance(statement.value, ast.Constant)
            data = namespace['load_excel_data'](ROOT / f'data_model_{task}.xlsx', task)
            self.assertEqual(len(data['test']), 160)
            with tempfile.TemporaryDirectory() as folder:
                run = namespace['NotebookRun'](task, data, ROOT / f'data_model_{task}.xlsx', Path(folder) / 'run')
                for method in ['measure_before', 'train', 'finish', 'export']:
                    with self.assertRaises(RuntimeError):
                        getattr(run, method)()

    def test_exported_bundles_reload_in_existing_backend(self):
        """Tiny random encoders for serialization only: no backward/optimizer/training."""
        import torch
        from transformers import AutoTokenizer, RobertaConfig, RobertaModel, RobertaForTokenClassification
        from .data import BIO
        from .notebook_runtime import NotebookRun
        from .inference import SOSPredictor
        tokenizer = AutoTokenizer.from_pretrained(ROOT.parent / 'nlp/sos_phobert_model', use_fast=False, local_files_only=True)
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / 'deployed'
            for task in ['a', 'b']:
                splits = load_excel_data(ROOT / f'data_model_{task}.xlsx', task)
                small = {k: v[:2] for k,v in splits.items()}
                config = RobertaConfig(vocab_size=len(tokenizer), hidden_size=8, num_hidden_layers=1,
                    num_attention_heads=2, intermediate_size=12, max_position_embeddings=258,
                    num_labels=len(BIO), id2label=dict(enumerate(BIO)), label2id={v:i for i,v in enumerate(BIO)})
                run = NotebookRun(task, small, ROOT / f'data_model_{task}.xlsx', Path(folder) / task)
                run.device = torch.device('cpu')
                constructor = 'AutoModel' if task == 'a' else 'AutoModelForTokenClassification'
                tiny = RobertaModel(config) if task == 'a' else RobertaForTokenClassification(config)
                with patch('ai.sos_ml.notebook_runtime.AutoTokenizer.from_pretrained', return_value=tokenizer), \
                     patch(f'ai.sos_ml.notebook_runtime.{constructor}.from_pretrained', return_value=tiny):
                    run.prepare()
                run.metadata['demo'] = True
                run.measure_before()
                run._save_best(run.before_validation, 0)
                # Supply a serialized checkpoint directly to test finish/export without train.
                run.training_done = True
                run.history = []
                comparison = run.finish()
                self.assertTrue(comparison)
                self.assertIn('text', run.predict('Cần nước uống tại xã An Bình.'))
                archive = run.export()
                with zipfile.ZipFile(archive) as z:
                    self.assertTrue(all(n.startswith(f'model_{task}/') for n in z.namelist()))
                    z.extractall(target)
                self.assertTrue((run.output / 'test_metrics.json').is_file())
            predictor = SOSPredictor(target, allow_demo=True)
            result = predictor.predict('Cần nước uống tại xã An Bình.')
            self.assertEqual(result['method'], 'phobert_two_model_v1')
            self.assertTrue(result['requires_review'])
            self.assertTrue(result['demo'])

    def test_train_selects_validation_and_preserves_parent_when_worse(self):
        """Exercise control flow with a scalar fake model, not PhoBERT training."""
        import torch
        from unittest.mock import Mock
        from .notebook_runtime import NotebookRun
        class ScalarModel(torch.nn.Module):
            def __init__(self):
                super().__init__()
                self.bias = torch.nn.Parameter(torch.zeros(5))
            def forward(self, input_ids, attention_mask):
                return self.bias[0].expand(len(input_ids)), self.bias[1:].expand(len(input_ids), 4)
        with tempfile.TemporaryDirectory() as folder:
            run = NotebookRun('a', {}, ROOT / 'data_model_a.xlsx', Path(folder) / 'run', epochs=2)
            run.output.mkdir()
            run.model, run.device = ScalarModel(), torch.device('cpu')
            run.baseline_done, run.parent = True, {'run_id':'parent'}
            run.before_validation = {'selection_score':.8}
            run.loaders = {'train':[{'input_ids':torch.ones((1,2), dtype=torch.long),
                'attention_mask':torch.ones((1,2), dtype=torch.long), 'targets':torch.ones((1,5))}],
                'validation': object()}
            run.evaluate = Mock(side_effect=[({'selection_score':.7}, []), ({'selection_score':.75}, [])])
            run._save_best = Mock()
            run.train()
            run._save_best.assert_called_once_with(run.before_validation, 0)
            self.assertTrue(run.training_done)
            self.assertEqual(len(run.history), 2)
            for call in run.evaluate.call_args_list:
                self.assertIs(call.args[1], run.loaders['validation'])


if __name__ == '__main__':
    unittest.main()
